"""
Slide ingestion.

Deliberate choice: render the deck to PDF once, then rasterize at whatever DPI
the target resolution needs. The Google Slides API `getThumbnail` endpoint caps
at 1600px wide, is an 'expensive read' (60 calls/min/user), and returns URLs
that expire in 30 minutes -- it cannot feed a 1920x1080 or 4K render.
Drive `files.export` -> PDF is ONE request per deck at unlimited resolution.
"""
from __future__ import annotations
from dataclasses import dataclass
from pathlib import Path
from realme.core.media import require, run, MediaError
from realme.pipeline import pdfdoc


def deck_to_pdf(src: Path, workdir: Path) -> Path:
    src = Path(src)
    if src.suffix.lower() == ".pdf":
        return src
    if src.suffix.lower() in {".pptx", ".ppt", ".odp"}:
        workdir.mkdir(parents=True, exist_ok=True)
        run([require("soffice"), "--headless", "--convert-to", "pdf",
             "--outdir", str(workdir), str(src)], f"LibreOffice convert {src.name}")
        out = workdir / (src.stem + ".pdf")
        if not out.exists():
            raise MediaError(f"LibreOffice produced no PDF for {src.name}.")
        return out
    raise MediaError(f"Unsupported deck format: {src.suffix}")


def deck_texts(deck: Path, workdir: Path) -> list[str]:
    """
    The text of every page of a deck, in page order.

    Exists because the revision path used to read this from `_work/deck.pdf` --
    a file nothing has ever written. The read was wrapped in a bare `except:
    pass` and guarded by `is_file()`, so it silently produced a list of empty
    strings, every slide matched nothing, and a lightly edited deck came back
    as forty brand-new slides. There is now exactly one way to ask a deck what
    it says, and it raises instead of shrugging.
    """
    deck, workdir = Path(deck), Path(workdir)
    if not deck.is_file():
        raise MediaError(f"The deck is not where it was recorded: {deck}")
    workdir.mkdir(parents=True, exist_ok=True)
    pdf = deck_to_pdf(deck, workdir)
    return [p.text for p in pdfdoc.read_pages(pdf)]


def rasterize(pdf: Path, outdir: Path, width_px: int = 1920) -> list[Path]:
    """Rasterize every page to PNG at the exact pixel width the video needs."""
    return pdfdoc.rasterize_pages(pdf, outdir, width_px)


@dataclass
class SlideContent:
    """What the script writer gets per slide."""
    index: int
    on_slide_text: str = ""
    speaker_notes: str = ""

    def as_prompt(self) -> str:
        parts = []
        if self.on_slide_text:
            parts.append(f"On-slide text:\n{self.on_slide_text}")
        if self.speaker_notes:
            # Notes come first in the model's attention for a reason: they are
            # what the instructor meant to say, whereas the on-slide text is
            # only what the audience can already read.
            parts.insert(0, f"INSTRUCTOR'S SPEAKER NOTES (follow these closely):\n"
                            f"{self.speaker_notes}")
        return "\n\n".join(parts) or "(no text on this slide)"


def extract_pdf_text(pdf: Path) -> list[str]:
    """Per-page visible text. Not speaker notes -- PDFs do not carry them."""
    try:
        return [pg.text for pg in pdfdoc.read_pages(pdf)]
    except Exception:
        return []


def _presentation(pptx: Path):
    """
    Open a .pptx, or say plainly why not.

    This used to `return []` on ImportError. That is the worst possible
    handling for this particular dependency: notes are the whole reason to
    prefer .pptx over .pdf, so losing them silently means the model narrates
    your bullet points back at the class and nothing anywhere says why.
    python-pptx is a hard dependency of RealMe, so a missing one is a broken
    install, not a configuration choice.
    """
    try:
        from pptx import Presentation
    except ImportError as e:
        raise MediaError(
            "python-pptx is not installed, so speaker notes cannot be read "
            "from this deck.\n"
            "  It is a required dependency -- something is wrong with the "
            "install. Repair it:\n"
            "    00_Windows\\1_Install.bat\n"
            "  or, from 5_Command_Prompt.bat:\n"
            "    pip install python-pptx") from e
    return Presentation(str(pptx))


def extract_pptx_notes(pptx: Path) -> list[str]:
    """
    Real speaker notes from a .pptx.

    This is the single most valuable grounding signal in the whole pipeline and
    v1 and v2 both threw it away, approximating it with the visible slide text.
    The visible text is what the audience can already read; the notes are what
    the instructor actually intended to say. Feeding the model the wrong one
    produces narration that recites the slide.
    """
    notes: list[str] = []
    for slide in _presentation(pptx).slides:
        text = ""
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            if frame is not None:
                text = (frame.text or "").strip()
        notes.append(text)
    return notes


def extract_pptx_text(pptx: Path) -> list[str]:
    """Visible text per slide, straight from the shapes."""
    out: list[str] = []
    for slide in _presentation(pptx).slides:
        chunks = [sh.text_frame.text.strip() for sh in slide.shapes
                  if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip()]
        out.append("\n".join(chunks))
    return out


def extract_content(src: Path, pdf: Path) -> list[SlideContent]:
    """
    Gather everything the script writer should see, from the best source.

    A .pptx carries real speaker notes, so read them from the original rather
    than from the PDF we rendered for rasterization. Google Slides decks should
    be pulled via `presentations.get` -> `slideProperties.notesPage` ->
    `notesProperties.speakerNotesObjectId`, which is one request per deck; the
    docs warn the shape may be absent on some slides, so handle empties.
    """
    src = Path(src)
    page_text = extract_pdf_text(pdf)
    notes: list[str] = []
    if src.suffix.lower() in {".pptx", ".ppt"}:
        notes = extract_pptx_notes(src)
        shape_text = extract_pptx_text(src)
        if shape_text and len(shape_text) >= len(page_text):
            page_text = shape_text
    n = max(len(page_text), len(notes))
    return [SlideContent(index=i,
                         on_slide_text=page_text[i] if i < len(page_text) else "",
                         speaker_notes=notes[i] if i < len(notes) else "")
            for i in range(n)]


def extract_notes(pdf: Path) -> list[str]:
    """Backwards-compatible shim: per-page visible text."""
    return extract_pdf_text(pdf)


def slide_pages(outdir) -> list[Path]:
    """
    The rendered slides, in PAGE order. **One place. Only this one.**

    `sorted(glob("slide-*.png"))` sorts lexicographically, and the files are
    named `slide-1.png` … `slide-12.png` without padding, so that order is

        1, 10, 11, 12, 2, 3, 4, 5, 6, 7, 8, 9

    Every deck of ten slides or more was scrambled from that point on, while
    decks of nine were perfect -- which is why this survived so long. The
    narration stayed in the right order because it never came from this list;
    only the pictures moved, so the symptom was "the notes are right and the
    slides are wrong".

    Sorting on the NUMBER, parsed from the name, is correct for both padded and
    unpadded files, so slides rasterised by an older version need no redo.
    """
    import re
    out = []
    for f in Path(outdir).glob("slide-*.png"):
        m = re.search(r"slide-(\d+)", f.stem)
        if m:
            out.append((int(m.group(1)), f))
    return [f for _, f in sorted(out)]
